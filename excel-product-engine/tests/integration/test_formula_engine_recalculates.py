"""Proves the formula engine's output isn't just well-formed text — it
actually calculates correctly, including structured table references,
SUMIFS with quoted criteria, and INDEX/MATCH, all recalculated by a
real spreadsheet engine (LibreOffice, via the project's recalc script).
"""

import json
import subprocess
import sys

from pathlib import Path

from excel_engine.core.workbook import ExcelWorkbook
from excel_engine.data.demo_data import write_demo_rows
from excel_engine.data.tables import ColumnSchema, add_typed_table
from excel_engine.formulas import Formula
from excel_engine.styles.style_manager import StyleManager

RECALC_SCRIPT = str(Path(__file__).resolve().parents[2] / "scripts" / "recalc.py")


def test_formula_engine_output_recalculates_correctly(tmp_path):
    style = StyleManager.for_theme("premium")
    workbook = ExcelWorkbook()

    orders = workbook.add_sheet("Orders")
    columns = [
        ColumnSchema(header="Customer", type="text"),
        ColumnSchema(header="Colorway", type="list", options=("Amber", "Teal")),
        ColumnSchema(header="Amount", type="currency"),
    ]
    add_typed_table(orders, style, columns, n_rows=10, table_name="Orders")
    write_demo_rows(
        orders, style, headers=["Customer", "Colorway", "Amount"],
        rows=[
            {"Customer": "Ananya Verma", "Colorway": "Amber", "Amount": 2499},
            {"Customer": "Rohan Gupta", "Colorway": "Teal", "Amount": 1899},
            {"Customer": "Priya Shah", "Colorway": "Amber", "Amount": 3200},
        ],
        top_left="A2",
    )

    dashboard = workbook.add_sheet("Dashboard")
    dashboard.set_formula("A1", Formula.sum("Orders", "Amount"))
    dashboard.set_formula("A2", Formula.average("Orders", "Amount"))
    dashboard.set_formula("A3", Formula.counta("Orders", "Customer"))
    dashboard.set_formula("A4", Formula.sumifs("Orders", "Amount", ("Colorway", "Amber")))
    dashboard.set_formula(
        "A5", Formula.index_match("Orders", "Amount", '"Rohan Gupta"', "Orders", "Customer")
    )
    dashboard.set_formula("A6", Formula.percentage_of_total("A4", "A1"))
    dashboard.set_formula("A7", Formula.growth("A4", "A2"))

    out_path = workbook.save(tmp_path / "formula_test.xlsx")

    result = subprocess.run(
        [sys.executable, RECALC_SCRIPT, str(out_path), "20"],
        capture_output=True, text=True, timeout=60,
    )
    report = json.loads(result.stdout)
    assert report["status"] == "success", report
    assert report["total_errors"] == 0
    assert report["total_formulas"] == 7

    import openpyxl
    reloaded = openpyxl.load_workbook(str(out_path), data_only=True)
    dash = reloaded["Dashboard"]
    assert dash["A1"].value == 7598          # SUM: 2499+1899+3200
    assert round(dash["A2"].value, 2) == round(7598 / 3, 2)  # AVERAGE
    assert dash["A3"].value == 3              # COUNTA
    assert dash["A4"].value == 5699           # SUMIFS Amber: 2499+3200
    assert dash["A5"].value == 1899           # INDEX/MATCH -> Rohan's amount
