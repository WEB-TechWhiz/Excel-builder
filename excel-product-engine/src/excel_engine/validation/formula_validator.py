"""Formula validation: formulas exist where expected, and every
structured table reference / sheet-qualified reference in every
formula in the workbook actually points at something real.

This is a regex-based scanner, not a full formula parser — it's scoped
to the two reference shapes this engine's own formula builders produce
(`Table[Column]` and `Sheet!range` / `'Sheet Name'!range`), which covers
the formulas worth validating in a workbook built with this engine.
"""

from __future__ import annotations

import re

from openpyxl.utils.cell import range_boundaries

from excel_engine.core.workbook import ExcelWorkbook
from excel_engine.exceptions.errors import CellAddressError
from excel_engine.validation.report import ValidationIssue, ValidationResult

_TABLE_REF_RE = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)\[([^\]]+)\]")
_SHEET_REF_RE = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!")


def _table_column_headers(raw_ws, table_ref: str) -> set[str]:  # type: ignore[no-untyped-def]
    try:
        min_col, min_row, max_col, _max_row = range_boundaries(table_ref)
    except (CellAddressError, ValueError, TypeError):
        return set()
    return {raw_ws.cell(row=min_row, column=col).value for col in range(min_col, max_col + 1)}


def _collect_tables(workbook: ExcelWorkbook) -> dict[str, set[str]]:
    all_tables: dict[str, set[str]] = {}
    for sheet_name in workbook.sheet_names:
        raw_ws = workbook.get_sheet(sheet_name).raw
        for table_name, table_ref in raw_ws.tables.items():
            all_tables[table_name] = _table_column_headers(raw_ws, table_ref)
    return all_tables


def validate_formulas(
    workbook: ExcelWorkbook,
    expected_formula_cells: list[tuple[str, str]] | None = None,
) -> ValidationResult:
    """`expected_formula_cells` is a list of (sheet, cell) pairs that
    must hold a formula (any formula — this checks presence, not the
    formula's own correctness beyond the reference checks below).
    """
    issues: list[ValidationIssue] = []
    valid_sheets = set(workbook.sheet_names)
    all_tables = _collect_tables(workbook)

    for sheet_name, cell_addr in expected_formula_cells or []:
        if not workbook.has_sheet(sheet_name):
            issues.append(ValidationIssue(
                category="formula",
                message=f"Expected a formula at {sheet_name}!{cell_addr}, but sheet "
                        f"{sheet_name!r} doesn't exist",
            ))
            continue
        value = workbook.get_sheet(sheet_name).get_value(cell_addr)
        if not (isinstance(value, str) and value.startswith("=")):
            issues.append(ValidationIssue(
                category="formula",
                message=f"Expected a formula at {sheet_name}!{cell_addr}, found: {value!r}",
            ))

    for sheet_name in workbook.sheet_names:
        raw_ws = workbook.get_sheet(sheet_name).raw
        for row in raw_ws.iter_rows():
            for cell in row:
                formula = cell.value
                if not (isinstance(formula, str) and formula.startswith("=")):
                    continue

                if "#REF!" in formula:
                    issues.append(ValidationIssue(
                        category="formula",
                        message=f"{sheet_name}!{cell.coordinate} contains a broken "
                                f"reference: {formula}",
                    ))

                for table_name, column_name in _TABLE_REF_RE.findall(formula):
                    if table_name not in all_tables:
                        issues.append(ValidationIssue(
                            category="formula",
                            message=f"{sheet_name}!{cell.coordinate} references table "
                                    f"{table_name!r}, which doesn't exist: {formula}",
                        ))
                    elif column_name not in all_tables[table_name]:
                        issues.append(ValidationIssue(
                            category="formula",
                            message=f"{sheet_name}!{cell.coordinate} references column "
                                    f"{column_name!r} on table {table_name!r}, which "
                                    f"doesn't have that column: {formula}",
                        ))

                for quoted, unquoted in _SHEET_REF_RE.findall(formula):
                    ref_sheet = quoted or unquoted
                    if ref_sheet not in valid_sheets:
                        issues.append(ValidationIssue(
                            category="formula",
                            message=f"{sheet_name}!{cell.coordinate} references sheet "
                                    f"{ref_sheet!r}, which doesn't exist: {formula}",
                        ))

    return ValidationResult(category="Formulas", issues=tuple(issues))
