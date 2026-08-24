"""Data table component — purely structural/visual: an Excel Table with
a themed header row and frozen panes. Column semantics (types, formulas,
dropdowns) are a Phase 5 (data engine) concern that will build on top
of this rather than duplicate it.
"""

from __future__ import annotations

import re

from openpyxl.worksheet.table import Table, TableStyleInfo

from excel_engine.core.cell import CellAddress
from excel_engine.core.range import CellRange
from excel_engine.core.worksheet import Worksheet
from excel_engine.exceptions.errors import WorkbookBuildError
from excel_engine.styles.style_manager import StyleManager


def _sanitize_table_name(name: str) -> str:
    sanitized = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not sanitized or sanitized[0].isdigit():
        sanitized = f"t_{sanitized}"
    return sanitized


def add_data_table(
    ws: Worksheet,
    style: StyleManager,
    headers: list[str],
    n_rows: int,
    table_name: str,
    top_left: str | CellAddress = "A1",
) -> CellRange:
    """Create a styled, banded Excel Table: a header row followed by
    `n_rows` blank (bordered, ready-to-fill) rows. Returns the full
    CellRange including the header row. Freezes panes below the header.
    """
    if not headers:
        raise WorkbookBuildError("add_data_table requires at least one header")
    if n_rows < 1:
        raise WorkbookBuildError("add_data_table requires n_rows >= 1")

    anchor = top_left if isinstance(top_left, CellAddress) else CellAddress.from_a1(top_left)

    for i, header in enumerate(headers):
        cell_addr = anchor.offset(columns=i)
        ws.set_value(cell_addr.to_a1(), header)
        cell = ws.raw[cell_addr.to_a1()]
        cell.font = style.header_font
        cell.fill = style.header_fill
        cell.alignment = style.center
        # Generic fallback so a bare header (no type info available here)
        # is never truncated — data.tables.add_typed_table overrides this
        # with better, type-aware widths once it knows column types.
        ws.set_column_width(cell_addr.column_letter, max(12, len(header) + 4))

    data_end = anchor.offset(rows=n_rows, columns=len(headers) - 1)
    full_range = CellRange(start=anchor, end=data_end)

    for row in range(anchor.row + 1, data_end.row + 1):
        for col in range(anchor.column, data_end.column + 1):
            ws.raw.cell(row=row, column=col).border = style.thin_border

    table = Table(displayName=_sanitize_table_name(table_name), ref=full_range.to_a1())
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    ws.raw.add_table(table)
    ws.freeze_panes(anchor.offset(rows=1).to_a1())

    return full_range
